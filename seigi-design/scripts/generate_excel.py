#!/usr/bin/env python3
"""せいぎデザイン Excel 帳票一括生成。"""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
THEME = json.loads((ROOT / "tokens" / "seigi.json").read_text(encoding="utf-8"))
OUT = ROOT / "excel" / "templates"

ORG = THEME["organization"]
DEPT = THEME["department"]
NAME = THEME["systemName"]
COLORS = THEME["colors"]

FILL_BRAND = PatternFill("solid", fgColor=COLORS["brand"].lstrip("#"))
FILL_LIGHT = PatternFill("solid", fgColor=COLORS["brandLight"].lstrip("#"))
FILL_ACCENT = PatternFill("solid", fgColor=COLORS["accent"].lstrip("#"))
FILL_SKY = PatternFill("solid", fgColor=COLORS["sky"].lstrip("#"))
FILL_SURFACE = PatternFill("solid", fgColor=COLORS["surface"].lstrip("#"))
FILL_WHITE = PatternFill("solid", fgColor="FFFFFF")

FONT_WHITE = Font(name="Yu Gothic", bold=True, color="FFFFFF", size=14)
FONT_WHITE_SM = Font(name="Yu Gothic", bold=True, color="FFFFFF", size=11)
FONT_TITLE = Font(name="Yu Gothic", bold=True, color=COLORS["brand"].lstrip("#"), size=18)
FONT_HEAD = Font(name="Yu Gothic", bold=True, color=COLORS["brand"].lstrip("#"), size=12)
FONT_BODY = Font(name="Yu Gothic", color=COLORS["ink"].lstrip("#"), size=11)
FONT_SOFT = Font(name="Yu Gothic", color=COLORS["inkSoft"].lstrip("#"), size=10)

THIN = Border(
    left=Side(style="thin", color=COLORS["muted"].lstrip("#")),
    right=Side(style="thin", color=COLORS["muted"].lstrip("#")),
    top=Side(style="thin", color=COLORS["muted"].lstrip("#")),
    bottom=Side(style="thin", color=COLORS["muted"].lstrip("#")),
)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def style_range(ws, cells, fill=None, font=None, alignment=None, border=True):
    for row in ws[cells]:
        for cell in row:
            if fill:
                cell.fill = fill
            if font:
                cell.font = font
            if alignment:
                cell.alignment = alignment
            if border:
                cell.border = THIN


def set_widths(ws, widths: dict[int, float]):
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width


def brand_banner(ws, title: str, subtitle: str, merge="A1:H2"):
    ws.merge_cells(merge)
    cell = ws[merge.split(":")[0]]
    cell.value = f"{title}\n{subtitle}"
    cell.fill = FILL_BRAND
    cell.font = FONT_WHITE
    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 22
    # accent strip
    ws.merge_cells("A3:H3")
    ws["A3"].value = f"{ORG}　{DEPT}　／　{NAME}"
    ws["A3"].fill = FILL_SKY
    ws["A3"].font = FONT_SOFT
    ws["A3"].alignment = LEFT


def save(wb: Workbook, name: str):
    OUT.mkdir(parents=True, exist_ok=True)
    path = OUT / name
    wb.save(path)
    return path


def make_minutes():
    wb = Workbook()
    ws = wb.active
    ws.title = "議事録"
    brand_banner(ws, "議事録・連絡票", "会議の決定と宿題を残す")
    headers = ["項目", "内容"]
    rows = [
        ("会議名", ""),
        ("日時", ""),
        ("場所", ""),
        ("出席者", ""),
        ("議題", ""),
        ("決定事項", ""),
        ("宿題（担当／期限）", ""),
        ("連絡事項", ""),
        ("次回", ""),
    ]
    ws["A5"] = headers[0]
    ws["B5"] = headers[1]
    style_range(ws, "A5:B5", FILL_LIGHT, FONT_WHITE_SM, CENTER)
    for i, (k, v) in enumerate(rows, start=6):
        ws[f"A{i}"] = k
        ws[f"B{i}"] = v
        ws[f"A{i}"].fill = FILL_SURFACE
        ws[f"A{i}"].font = FONT_HEAD
        ws[f"B{i}"].font = FONT_BODY
        ws[f"A{i}"].border = THIN
        ws[f"B{i}"].border = THIN
        ws.row_dimensions[i].height = 28
    set_widths(ws, {1: 22, 2: 70})
    return save(wb, "seigi-minutes.xlsx")


def make_checklist():
    wb = Workbook()
    ws = wb.active
    ws.title = "チェックリスト"
    brand_banner(ws, "チェックリスト／点検表", "実施前に確認する項目")
    headers = ["No", "確認項目", "基準", "結果", "指摘", "確認者", "日付"]
    for i, h in enumerate(headers, start=1):
        cell = ws.cell(5, i, h)
        cell.fill = FILL_LIGHT
        cell.font = FONT_WHITE_SM
        cell.alignment = CENTER
        cell.border = THIN
    for r in range(6, 21):
        ws.cell(r, 1, r - 5).alignment = CENTER
        for c in range(1, 8):
            ws.cell(r, c).border = THIN
            ws.cell(r, c).font = FONT_BODY
            if r % 2 == 0:
                ws.cell(r, c).fill = FILL_SURFACE
        ws.cell(r, 4, "□ OK　□ NG")
    set_widths(ws, {1: 6, 2: 32, 3: 18, 4: 16, 5: 22, 6: 12, 7: 12})
    return save(wb, "seigi-checklist.xlsx")


def make_schedule():
    wb = Workbook()
    ws = wb.active
    ws.title = "工程表"
    brand_banner(ws, "工程表・簡易スケジュール", "週次の見える化")
    headers = ["WBS", "作業", "担当", "開始", "終了", "進捗%", "状態", "メモ"]
    for i, h in enumerate(headers, start=1):
        cell = ws.cell(5, i, h)
        cell.fill = FILL_BRAND
        cell.font = FONT_WHITE_SM
        cell.alignment = CENTER
        cell.border = THIN
    samples = [
        ("1", "要求整理", "生産技術", "", "", "0", "未着手", ""),
        ("2", "構想設計", "設計", "", "", "0", "未着手", ""),
        ("3", "ジグDR", "全員", "", "", "0", "予定", ""),
        ("4", "試作", "製作", "", "", "0", "予定", ""),
        ("5", "量産準備", "生産技術", "", "", "0", "予定", ""),
    ]
    for r, row in enumerate(samples, start=6):
        for c, val in enumerate(row, start=1):
            cell = ws.cell(r, c, val)
            cell.border = THIN
            cell.font = FONT_BODY
            cell.alignment = CENTER if c != 8 else LEFT
            if r % 2 == 0:
                cell.fill = FILL_SURFACE
    for r in range(11, 21):
        for c in range(1, 9):
            ws.cell(r, c).border = THIN
    set_widths(ws, {1: 8, 2: 22, 3: 12, 4: 12, 5: 12, 6: 10, 7: 12, 8: 28})
    return save(wb, "seigi-schedule.xlsx")


def make_metrics():
    wb = Workbook()
    cover = wb.active
    cover.title = "表紙"
    brand_banner(cover, "数値・集計シート", "表紙付きの報告用ブック", "A1:F2")
    cover["A5"] = "資料名"
    cover["B5"] = ""
    cover["A6"] = "対象期間"
    cover["B6"] = ""
    cover["A7"] = "作成者"
    cover["B7"] = ""
    cover["A8"] = "承認者"
    cover["B8"] = ""
    for r in range(5, 9):
        cover[f"A{r}"].fill = FILL_SURFACE
        cover[f"A{r}"].font = FONT_HEAD
        cover[f"A{r}"].border = THIN
        cover[f"B{r}"].border = THIN
        cover[f"B{r}"].font = FONT_BODY
    cover["A10"] = "ポイント（やわらかく・はっきり）"
    cover["A10"].font = FONT_TITLE
    cover["A11"] = "結論を先に書く。数字は単位付きで。"
    cover["A11"].font = FONT_BODY
    set_widths(cover, {1: 18, 2: 40, 3: 12, 4: 12, 5: 12, 6: 12})

    data = wb.create_sheet("集計")
    brand_banner(data, "集計", "明細")
    headers = ["区分", "項目", "計画", "実績", "差異", "コメント"]
    for i, h in enumerate(headers, start=1):
        cell = data.cell(5, i, h)
        cell.fill = FILL_LIGHT
        cell.font = FONT_WHITE_SM
        cell.alignment = CENTER
        cell.border = THIN
    for r in range(6, 21):
        for c in range(1, 7):
            cell = data.cell(r, c, "")
            cell.border = THIN
            cell.font = FONT_BODY
            if r % 2 == 0:
                cell.fill = FILL_SURFACE
    set_widths(data, {1: 14, 2: 24, 3: 12, 4: 12, 5: 12, 6: 30})
    return save(wb, "seigi-metrics.xlsx")


def make_jig_dr():
    wb = Workbook()
    ws = wb.active
    ws.title = "ジグDR指摘"
    brand_banner(ws, "ジグ設計レビュー（DR）帳票", "指摘・宿題管理")
    meta = [("案件名", "B5"), ("ワーク／工程", "B6"), ("レビュー日", "B7"), ("主催", "B8"), ("判定", "B9")]
    labels = ["案件名", "ワーク／工程", "レビュー日", "主催", "判定（合格／条件付／差戻し）"]
    for i, label in enumerate(labels, start=5):
        ws[f"A{i}"] = label
        ws[f"A{i}"].fill = FILL_SURFACE
        ws[f"A{i}"].font = FONT_HEAD
        ws[f"A{i}"].border = THIN
        ws[f"B{i}"].border = THIN
        ws.merge_cells(f"B{i}:H{i}")

    ws["A11"] = "指摘一覧"
    ws["A11"].font = FONT_TITLE
    headers = ["ID", "箇所", "内容", "重要度", "担当", "期限", "状態", "備考"]
    for i, h in enumerate(headers, start=1):
        cell = ws.cell(12, i, h)
        cell.fill = FILL_ACCENT if h in {"重要度", "状態"} else FILL_BRAND
        cell.font = FONT_WHITE_SM
        cell.alignment = CENTER
        cell.border = THIN
    for r in range(13, 28):
        ws.cell(r, 1, f"DR-{r-12:02d}")
        for c in range(1, 9):
            cell = ws.cell(r, c)
            cell.border = THIN
            cell.font = FONT_BODY
            if r % 2 == 0:
                cell.fill = FILL_SURFACE
        ws.cell(r, 7, "未／進行／完了")

    ws2 = wb.create_sheet("宿題")
    brand_banner(ws2, "宿題・期限", "フォロー用")
    h2 = ["ID", "宿題", "担当", "期限", "状態", "確認者", "完了日", "メモ"]
    for i, h in enumerate(h2, start=1):
        cell = ws2.cell(5, i, h)
        cell.fill = FILL_BRAND
        cell.font = FONT_WHITE_SM
        cell.alignment = CENTER
        cell.border = THIN
    for r in range(6, 21):
        for c in range(1, 9):
            cell = ws2.cell(r, c)
            cell.border = THIN
            cell.font = FONT_BODY
            if r % 2 == 0:
                cell.fill = FILL_SURFACE
    set_widths(ws, {1: 10, 2: 16, 3: 36, 4: 10, 5: 12, 6: 12, 7: 12, 8: 20})
    set_widths(ws2, {1: 10, 2: 28, 3: 12, 4: 12, 5: 12, 6: 12, 7: 12, 8: 22})
    return save(wb, "seigi-jig-dr.xlsx")


def make_department_text():
    wb = Workbook()
    ws = wb.active
    ws.title = "部門テキスト"
    brand_banner(ws, "部門テキスト差し込み例", "社名・部署・システム名の組み合わせ")
    ws["A5"] = "用途"
    ws["B5"] = "推奨表記"
    ws["C5"] = "メモ"
    style_range(ws, "A5:C5", FILL_BRAND, FONT_WHITE_SM, CENTER)
    samples = [
        ("資料ヘッダー", f"{DEPT} · {NAME}", "短く"),
        ("表紙サブ", f"{ORG}", "正式社名"),
        ("フッター", f"{ORG} / {DEPT}", "全資料共通"),
        ("ファイル名例", "YYYYMMDD_案件_せいぎ.pptx", "日付先頭"),
        ("DRタイトル", "ジグ設計レビュー（DR）", "案件名を前に付けても可"),
        ("注意ラベル", "注意（ノーリツレッド）", "強調のみ赤を使う"),
    ]
    for i, (a, b, c) in enumerate(samples, start=6):
        ws[f"A{i}"] = a
        ws[f"B{i}"] = b
        ws[f"C{i}"] = c
        for col in "ABC":
            ws[f"{col}{i}"].border = THIN
            ws[f"{col}{i}"].font = FONT_BODY
            if i % 2 == 0:
                ws[f"{col}{i}"].fill = FILL_SURFACE
    ws["A13"] = "マーク"
    ws["A13"].font = FONT_HEAD
    ws["B13"] = "assets/mark/seigi-mark.svg（雲＋歯車）。公式ロゴは使わない。"
    ws["B13"].font = FONT_BODY
    set_widths(ws, {1: 18, 2: 42, 3: 28})
    return save(wb, "seigi-department-text.xlsx")


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    for fn in [
        make_minutes,
        make_checklist,
        make_schedule,
        make_metrics,
        make_jig_dr,
        make_department_text,
    ]:
        path = fn()
        print(f"  {path.name}")
    print(f"TOTAL excel: {len(list(OUT.glob('*.xlsx')))}")


if __name__ == "__main__":
    main()
