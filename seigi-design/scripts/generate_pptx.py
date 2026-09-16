#!/usr/bin/env python3
"""せいぎデザイン PowerPoint 一括生成。"""

from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt, Emu

ROOT = Path(__file__).resolve().parents[1]
THEME = json.loads((ROOT / "tokens" / "seigi.json").read_text(encoding="utf-8"))
OUT = ROOT / "ppt" / "templates"
LAYOUTS = OUT / "layouts"
W, H = Inches(13.333), Inches(7.5)


def hex_rgb(value: str) -> RGBColor:
    value = value.lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


C = {k: hex_rgb(v) for k, v in THEME["colors"].items()}
ORG = THEME["organization"]
DEPT = THEME["department"]
NAME = THEME["systemName"]


def blank_prs() -> Presentation:
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    return prs


def add_blank(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])


def shape_fill(shape, color: RGBColor):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def oval(slide, l, t, w, h, color):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, l, t, w, h)
    shape_fill(s, color)
    return s


def round_rect(slide, l, t, w, h, color, adj=0.3):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    shape_fill(s, color)
    try:
        s.adjustments[0] = adj
    except Exception:
        pass
    return s


def rect(slide, l, t, w, h, color):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    shape_fill(s, color)
    return s


def textbox(slide, l, t, w, h, text, size, color, bold=False, align=PP_ALIGN.LEFT, font="Yu Gothic"):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    return box


def bullets(slide, l, t, w, h, lines, size=18, color=None):
    color = color or C["ink"]
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.level = 0
        run = p.add_run()
        run.text = f"• {line}"
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.name = "Yu Gothic"
        p.space_after = Pt(8)
    return box


def clouds(slide):
    oval(slide, Inches(-0.4), Inches(5.8), Inches(3.2), Inches(2.0), C["skyDeep"])
    oval(slide, Inches(1.2), Inches(6.1), Inches(3.8), Inches(2.2), C["cloud"])
    oval(slide, Inches(10.2), Inches(-0.6), Inches(3.5), Inches(2.2), C["skyDeep"])
    oval(slide, Inches(11.0), Inches(-0.2), Inches(3.0), Inches(1.8), C["surface"])


def header_bar(slide, title: str, subtitle: str | None = None):
    rect(slide, Inches(0), Inches(0), W, Inches(0.85), C["brand"])
    textbox(slide, Inches(0.5), Inches(0.18), Inches(9), Inches(0.5), title, 22, C["white"], bold=True)
    textbox(
        slide,
        Inches(9.5),
        Inches(0.22),
        Inches(3.4),
        Inches(0.45),
        f"{DEPT} · {NAME}",
        12,
        C["sky"],
        align=PP_ALIGN.RIGHT,
    )
    if subtitle:
        textbox(slide, Inches(0.5), Inches(1.0), Inches(12), Inches(0.4), subtitle, 14, C["inkSoft"])


def footer(slide, page: str = ""):
    textbox(slide, Inches(0.5), Inches(7.05), Inches(9), Inches(0.3), f"{ORG} / {DEPT}", 11, C["inkSoft"])
    textbox(slide, Inches(10.5), Inches(7.05), Inches(2.3), Inches(0.3), page, 11, C["inkSoft"], align=PP_ALIGN.RIGHT)
    rect(slide, Inches(0.5), Inches(6.95), Inches(12.3), Emu(20000), C["muted"])


def accent_chip(slide, l, t, label: str):
    round_rect(slide, l, t, Inches(1.6), Inches(0.38), C["accent"], 0.5)
    textbox(slide, l, t + Inches(0.04), Inches(1.6), Inches(0.32), label, 12, C["white"], bold=True, align=PP_ALIGN.CENTER)


def panel(slide, l, t, w, h, title: str, lines: list[str]):
    round_rect(slide, l, t, w, h, C["cloud"], 0.18)
    textbox(slide, l + Inches(0.25), t + Inches(0.2), w - Inches(0.4), Inches(0.4), title, 16, C["brand"], bold=True)
    bullets(slide, l + Inches(0.25), t + Inches(0.65), w - Inches(0.45), h - Inches(0.9), lines, size=14)


def decorate(slide):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = C["sky"]
    clouds(slide)


def slide_cover(prs, title: str, subtitle: str, tag: str = "テンプレート"):
    s = add_blank(prs)
    decorate(s)
    round_rect(s, Inches(0.8), Inches(1.6), Inches(11.7), Inches(4.2), C["cloud"], 0.12)
    accent_chip(s, Inches(1.1), Inches(1.95), tag)
    textbox(s, Inches(1.1), Inches(2.5), Inches(10.8), Inches(1.2), title, 36, C["brand"], bold=True)
    textbox(s, Inches(1.1), Inches(3.8), Inches(10.8), Inches(0.8), subtitle, 18, C["inkSoft"])
    textbox(s, Inches(1.1), Inches(5.0), Inches(10.8), Inches(0.4), f"{ORG}　{DEPT}", 14, C["brandLight"])
    footer(s, NAME)
    return s


def slide_section(prs, title: str, blurb: str, page=""):
    s = add_blank(prs)
    decorate(s)
    header_bar(s, title)
    round_rect(s, Inches(0.8), Inches(2.0), Inches(11.7), Inches(3.2), C["cloud"], 0.12)
    textbox(s, Inches(1.2), Inches(2.6), Inches(10.8), Inches(1.8), blurb, 24, C["ink"])
    footer(s, page)
    return s


def slide_bullets(prs, title: str, lines: list[str], page=""):
    s = add_blank(prs)
    decorate(s)
    header_bar(s, title)
    round_rect(s, Inches(0.7), Inches(1.3), Inches(11.9), Inches(5.2), C["cloud"], 0.12)
    bullets(s, Inches(1.1), Inches(1.7), Inches(11), Inches(4.5), lines, size=20)
    footer(s, page)
    return s


def slide_two_col(prs, title: str, left_t: str, left: list[str], right_t: str, right: list[str], page=""):
    s = add_blank(prs)
    decorate(s)
    header_bar(s, title)
    panel(s, Inches(0.7), Inches(1.3), Inches(5.7), Inches(5.2), left_t, left)
    panel(s, Inches(6.8), Inches(1.3), Inches(5.7), Inches(5.2), right_t, right)
    footer(s, page)
    return s


def slide_tableish(prs, title: str, headers: list[str], rows: list[list[str]], page=""):
    s = add_blank(prs)
    decorate(s)
    header_bar(s, title)
    cols = len(headers)
    total_w = Inches(12.0)
    col_w = total_w / cols
    top = Inches(1.5)
    left0 = Inches(0.65)
    row_h = Inches(0.55)
    for i, h in enumerate(headers):
        round_rect(s, left0 + col_w * i, top, col_w - Inches(0.08), row_h, C["brand"], 0.2)
        textbox(s, left0 + col_w * i, top + Inches(0.1), col_w - Inches(0.08), row_h, h, 13, C["white"], bold=True, align=PP_ALIGN.CENTER)
    for r, row in enumerate(rows):
        y = top + row_h * (r + 1) + Inches(0.08) * (r + 1)
        bg = C["cloud"] if r % 2 == 0 else C["surface"]
        for i, cell in enumerate(row):
            round_rect(s, left0 + col_w * i, y, col_w - Inches(0.08), row_h, bg, 0.2)
            textbox(s, left0 + col_w * i + Inches(0.08), y + Inches(0.1), col_w - Inches(0.2), row_h, cell, 12, C["ink"], align=PP_ALIGN.CENTER)
    footer(s, page)
    return s


def slide_kpi(prs, title: str, kpis: list[tuple[str, str]], page=""):
    s = add_blank(prs)
    decorate(s)
    header_bar(s, title)
    n = len(kpis)
    width = Inches(11.8) / n
    for i, (label, value) in enumerate(kpis):
        x = Inches(0.7) + width * i
        round_rect(s, x, Inches(2.2), width - Inches(0.2), Inches(2.8), C["cloud"], 0.15)
        textbox(s, x + Inches(0.2), Inches(2.5), width - Inches(0.5), Inches(0.5), label, 14, C["inkSoft"], align=PP_ALIGN.CENTER)
        textbox(s, x + Inches(0.2), Inches(3.3), width - Inches(0.5), Inches(1.0), value, 32, C["brand"], bold=True, align=PP_ALIGN.CENTER)
    footer(s, page)
    return s


def save(prs: Presentation, name: str):
    OUT.mkdir(parents=True, exist_ok=True)
    path = OUT / name
    prs.save(path)
    return path, len(prs.slides)


def build_basic() -> tuple[str, int]:
    prs = blank_prs()
    slide_cover(prs, "せいぎデザイン 基本セット", "表紙・アジェンダ・締めまで揃えるスターター", "基本")
    slide_bullets(prs, "アジェンダ", ["本日の目的", "現状共有", "論点整理", "次アクション", "クロージング"], "2")
    slide_section(prs, "本日の目的", "参加者が同じゴールを持てるよう、1文で目的を書く", "3")
    slide_two_col(prs, "ポイント整理", "伝えること", ["背景", "結論先出し", "根拠"], "確認すること", ["前提", "制約", "期限"], "4")
    slide_bullets(prs, "クロージング", ["決定事項の確認", "宿題と担当", "次回予定"], "5")
    slide_cover(prs, "ご清聴ありがとうございました", f"{ORG} / {DEPT}", "Thanks")
    return save(prs, "seigi-basic.pptx")


def build_kaizen() -> tuple[str, int]:
    prs = blank_prs()
    slide_cover(prs, "改善報告テンプレート", "現状 → 問題 → なぜなぜ → 対策 → 効果 → 横展", "改善")
    slide_bullets(prs, "現状", ["対象工程 / 設備", "現状の数値・事実", "関係者"], "2")
    slide_bullets(prs, "問題", ["困っていること（事実）", "影響（安全・品質・コスト・納期）", "放置した場合のリスク"], "3")
    slide_two_col(prs, "なぜなぜ", "なぜ1〜3", ["なぜ起きたか", "なぜ検出できなかったか", "なぜ再発しうるか"], "真因候補", ["人", "設備", "方法", "材料"], "4")
    slide_bullets(prs, "対策", ["暫定対策", "恒久対策", "標準化・教育", "効果確認方法"], "5")
    slide_kpi(prs, "効果", [("不良率", "−40%"), ("段取り", "−12分"), ("横展", "3工程")], "6")
    slide_bullets(prs, "横展", ["適用候補ライン", "必要な治具・手順書", "スケジュール"], "7")
    return save(prs, "seigi-kaizen.pptx")


def build_process() -> tuple[str, int]:
    prs = blank_prs()
    slide_cover(prs, "工程・設備説明", "流れとポイントをやわらかく・はっきり", "工程")
    slide_bullets(prs, "工程フロー", ["投入", "加工", "検査", "出荷／次工程"], "2")
    slide_two_col(prs, "設備概要", "仕様", ["型式", "能力", "設置場所"], "ポイント", ["安全", "品質", "保全"], "3")
    slide_bullets(prs, "説明のコツ", ["写真＋一言注釈", "数字は大きく", "注意は赤アクセント"], "4")
    return save(prs, "seigi-process.pptx")


def build_standup() -> tuple[str, int]:
    prs = blank_prs()
    slide_cover(prs, "朝会・短い共有", "3〜5枚で要点だけ", "朝会")
    slide_kpi(prs, "本日の指標", [("安全", "0件"), ("品質", "良品率"), ("進捗", "予定比")], "2")
    slide_bullets(prs, "共有事項", ["昨日の完了", "今日の重点", "困りごと"], "3")
    slide_bullets(prs, "依頼・確認", ["応援が必要な作業", "部品・治具の遅れ", "安全注意"], "4")
    return save(prs, "seigi-standup.pptx")


def build_jig_dr() -> tuple[str, int]:
    prs = blank_prs()
    slide_cover(prs, "ジグ設計レビュー（DR）", "設計審査＋デザインレビューを同じ型で", "ジグDR")
    slide_bullets(prs, "目的・要求", ["ワーク／工程", "ねらい（精度・タクト・安全）", "制約（スペース・コスト・期限）"], "2")
    slide_two_col(prs, "構造・機構", "構造", ["固定方法", " Locating", "クランプ"], "機構", ["駆動", "センサ", "可動範囲"], "3")
    slide_bullets(prs, "安全・干渉", ["挟まれ・切れ", "ポカヨケ", "干渉チェック（人・設備・周辺）"], "4")
    slide_bullets(prs, "公差・材質", ["重要寸法", "公差根拠", "材質／表面処理"], "5")
    slide_tableish(
        prs,
        "指摘一覧",
        ["ID", "箇所", "内容", "重要度", "担当"],
        [
            ["DR-01", "クランプ", "例: 開き量不足", "高", "設計"],
            ["DR-02", "ガイド", "例: 摩耗対策", "中", "生産技術"],
            ["DR-03", "表示", "例: 向き明示", "低", "設計"],
        ],
        "6",
    )
    slide_tableish(
        prs,
        "宿題・期限",
        ["ID", "宿題", "担当", "期限", "状態"],
        [
            ["A-01", "干渉図更新", "〇〇", "MM/DD", "未"],
            ["A-02", "強度確認", "〇〇", "MM/DD", "進行"],
            ["A-03", "試作計画", "〇〇", "MM/DD", "完了"],
        ],
        "7",
    )
    slide_two_col(prs, "合否・見解", "判定", ["合格", "条件付き合格", "差戻し"], "条件・コメント", ["必須宿題", "再レビュー日", "残リスク"], "8")
    return save(prs, "seigi-jig-dr.pptx")


def build_generic() -> tuple[str, int]:
    prs = blank_prs()
    slide_cover(prs, "汎用スライド集", "比較・引用・写真注釈・注意・Q&A", "汎用")
    slide_two_col(prs, "2列比較", "Before", ["課題A", "課題B"], "After", ["改善A", "改善B"], "2")
    slide_section(prs, "引用・メッセージ", "「結論を一文で置く」", "3")
    slide_bullets(prs, "写真＋注釈（置き場）", ["左に写真、右に注釈", "赤で危険箇所、青で要点", "1スライド1メッセージ"], "4")
    slide_tableish(
        prs,
        "スケジュール",
        ["週", "内容", "担当", "成果物"],
        [["W1", "要件固め", "PT", "要求書"], ["W2", "構想", "設計", "構想図"], ["W3", "DR", "全員", "指摘表"]],
        "5",
    )
    slide_bullets(prs, "注意喚起", ["安全第一", "仮置き禁止", "手順遵守"], "6")
    slide_two_col(prs, "Q&A", "よくある質問", ["納期は？", "試作は何台？"], "回答メモ", ["根拠を書く", "代替案も用意"], "7")
    return save(prs, "seigi-generic.pptx")


def build_starter() -> tuple[str, int]:
    prs = blank_prs()
    slide_cover(prs, "せいぎデザイン スターター", "まずはこのデッキから始める", "Starter")
    slide_bullets(prs, "このキットに含むもの", ["基本セット", "改善報告", "ジグDR", "汎用部品", "Excel帳票は別ZIP"], "2")
    slide_bullets(prs, "使い方", ["表紙のタイトルを書き換える", "不要スライドを削除", "色は変えなくてOK（トークン準拠）"], "3")
    slide_cover(prs, "統一は意識しなくてよい", "テンプレを開いた瞬間に、せいぎデザインになっている", "Tip")
    return save(prs, "seigi-starter.pptx")


def build_layouts():
    LAYOUTS.mkdir(parents=True, exist_ok=True)
    specs = [
        ("cover", lambda prs: slide_cover(prs, "タイトルを入れる", "サブタイトル", "表紙")),
        ("section", lambda prs: slide_section(prs, "節タイトル", "リード文をここに")),
        ("bullets", lambda prs: slide_bullets(prs, "箇条書き", ["項目1", "項目2", "項目3"])),
        ("two-col", lambda prs: slide_two_col(prs, "2列", "左", ["A", "B"], "右", ["C", "D"])),
        ("kpi", lambda prs: slide_kpi(prs, "KPI", [("指標1", "00"), ("指標2", "00"), ("指標3", "00")])),
        ("before-after", lambda prs: slide_two_col(prs, "Before / After", "Before", ["現状"], "After", ["あるべき姿"])),
        ("risk", lambda prs: slide_bullets(prs, "リスク", ["重大度", "発生頻度", "検知性", "対策"])),
        ("checklist", lambda prs: slide_bullets(prs, "チェック", ["□ 項目1", "□ 項目2", "□ 項目3"])),
        ("photo-notes", lambda prs: slide_two_col(prs, "写真グリッド", "写真エリア", ["画像を配置"], "注釈", ["ポイント1", "ポイント2"])),
        ("callout", lambda prs: slide_section(prs, "注釈吹き出し", "ここがポイント、を大きく")),
        ("agenda", lambda prs: slide_bullets(prs, "アジェンダ", ["1.", "2.", "3.", "4."])),
        ("thanks", lambda prs: slide_cover(prs, "ありがとうございました", f"{DEPT}", "Thanks")),
        (
            "header-footer",
            lambda prs: slide_bullets(
                prs, "ヘッダー／フッター例", ["上部に部署名", "下部に会社名", "赤は注意のみ"]
            ),
        ),
    ]
    count = 0
    for name, builder in specs:
        prs = blank_prs()
        builder(prs)
        prs.save(LAYOUTS / f"{name}.pptx")
        count += 1
    return count


def build_all_in_one() -> tuple[str, int]:
    prs = blank_prs()
    slide_cover(prs, "せいぎデザイン 総合デッキ", "基本・改善・工程・朝会・ジグDR・汎用を1本に", "総合")
    for title, lines in [
        ("基本", ["目的共有", "アジェンダ", "クロージング"]),
        ("改善", ["現状→対策→効果→横展"]),
        ("工程・設備", ["フロー", "仕様", "安全"]),
        ("朝会", ["指標", "共有", "依頼"]),
        ("ジグDR", ["要求", "構造", "安全", "指摘", "宿題", "合否"]),
        ("汎用部品", ["比較", "KPI", "表", "注意", "Q&A"]),
    ]:
        slide_bullets(prs, title, lines)
    return save(prs, "seigi-all-themes-catalog.pptx")


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    results = []
    for builder in [
        build_basic,
        build_kaizen,
        build_process,
        build_standup,
        build_jig_dr,
        build_generic,
        build_starter,
        build_all_in_one,
    ]:
        path, n = builder()
        results.append((path.name, n))
        print(f"  {path.name}: {n} slides")
    n_layouts = build_layouts()
    print(f"  layouts/: {n_layouts} files")
    print(f"TOTAL decks: {len(results)} + layouts {n_layouts}")


if __name__ == "__main__":
    main()
