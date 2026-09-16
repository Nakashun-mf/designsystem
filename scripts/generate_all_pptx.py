#!/usr/bin/env python3
"""Generate soft-cloud-first PowerPoint templates with fluffy oval clouds."""

from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
THEMES_DIR = ROOT / "tokens" / "themes"
OUT_DIR = ROOT / "ppt" / "templates"
W, H = Inches(13.333), Inches(7.5)


def hex_rgb(value: str) -> RGBColor:
    value = value.lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def load_themes() -> dict:
    themes = {}
    for path in sorted(THEMES_DIR.glob("*.json")):
        themes[path.stem] = json.loads(path.read_text(encoding="utf-8"))
    # soft-cloud first
    if "soft-cloud" in themes:
        ordered = {"soft-cloud": themes["soft-cloud"]}
        for k, v in themes.items():
            if k != "soft-cloud":
                ordered[k] = v
        return ordered
    return themes


def blank_prs() -> Presentation:
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    return prs


def add_blank(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])


def oval(slide, left, top, width, height, fill: RGBColor):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    return shape


def round_rect(slide, left, top, width, height, fill: RGBColor, adj: float = 0.35):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    try:
        shape.adjustments[0] = adj
    except Exception:
        pass
    return shape


def rect(slide, left, top, width, height, fill: RGBColor):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    return shape


def textbox(slide, left, top, width, height, text: str, size: int, color: RGBColor,
            bold: bool = False, font: str = "Yu Gothic", align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    return box


def c(theme: dict, role: str, fallback: str | None = None) -> RGBColor:
    colors = theme["colors"]
    if role in colors:
        return hex_rgb(colors[role])
    if fallback and fallback in colors:
        return hex_rgb(colors[fallback])
    # soft defaults
    defaults = {
        "sky": "#d9efff", "cloud": "#ffffff", "blush": "#ffd0dc",
        "peach": "#ffe3d4", "surface": "#fff0f4", "paper": "#fff7fa",
    }
    return hex_rgb(defaults.get(role, colors.get("brand", "#f07a9a")))


def bg(theme):
    return c(theme, theme.get("pptBgRole", "sky"), "sky")


def fg(theme):
    return c(theme, theme.get("pptFgRole", "ink"), "ink")


def panel(theme):
    return c(theme, theme.get("pptPanelRole", "cloud"), "cloud")


def pfg(theme):
    return c(theme, theme.get("pptPanelFgRole", "ink"), "ink")


def brand(theme):
    return c(theme, "brand")


def on_brand(theme):
    if theme["id"] == "wishing-umbrella":
        return c(theme, "paper")
    return c(theme, "white", "cloud")


def draw_cloud(slide, x, y, scale=1.0, fill: RGBColor | None = None, theme=None):
    """Fluffy cloud from overlapping ovals. x,y in inches for the left of the base."""
    fill = fill or (c(theme, "cloud") if theme else RGBColor(255, 255, 255))
    # base
    oval(slide, Inches(x), Inches(y), Inches(2.6 * scale), Inches(1.15 * scale), fill)
    oval(slide, Inches(x + 0.35 * scale), Inches(y - 0.7 * scale), Inches(1.5 * scale), Inches(1.4 * scale), fill)
    oval(slide, Inches(x + 1.15 * scale), Inches(y - 0.55 * scale), Inches(1.35 * scale), Inches(1.25 * scale), fill)
    oval(slide, Inches(x + 1.9 * scale), Inches(y - 0.25 * scale), Inches(1.0 * scale), Inches(1.0 * scale), fill)


def draw_bokeh(slide, theme, points=None):
    points = points or [(2.2, 1.5, 0.28), (11.2, 1.2, 0.36), (6.5, 6.4, 0.22), (10.4, 5.8, 0.3), (3.8, 5.2, 0.18)]
    for i, (x, y, r) in enumerate(points):
        fill = c(theme, "bokeh") if i % 2 == 0 else c(theme, "gold", "peach")
        oval(slide, Inches(x), Inches(y), Inches(r * 2), Inches(r * 2), fill)


def draw_umbrella(slide, x, y, scale=1.0, theme=None):
    """Simple cute umbrella: blue canopy + stick."""
    blue = c(theme, "umbrella") if theme else hex_rgb("#3f7fd4")
    soft = c(theme, "umbrellaSoft", "umbrella") if theme else hex_rgb("#6ea0e6")
    ink = c(theme, "inkSoft", "ink") if theme else hex_rgb("#8a6574")
    # canopy (rounded top)
    oval(slide, Inches(x), Inches(y), Inches(1.8 * scale), Inches(1.15 * scale), soft)
    round_rect(slide, Inches(x), Inches(y + 0.45 * scale), Inches(1.8 * scale), Inches(0.45 * scale), blue, 0.2)
    # stick
    rect(slide, Inches(x + 0.85 * scale), Inches(y + 0.85 * scale), Inches(0.08 * scale), Inches(0.85 * scale), ink)
    # handle curve approximation
    oval(slide, Inches(x + 0.78 * scale), Inches(y + 1.55 * scale), Inches(0.28 * scale), Inches(0.28 * scale), ink)


def draw_poms(slide, x, y, theme, n=5):
    colors = ["brand", "umbrella", "gold", "lavender", "peach"]
    for i in range(n):
        oval(slide, Inches(x + i * 0.38), Inches(y), Inches(0.28), Inches(0.28), c(theme, colors[i % len(colors)]))


def sky_backdrop(slide, theme):
    rect(slide, 0, 0, W, H, bg(theme))
    # warm peach / lavender bands (MV-like soft glow)
    round_rect(slide, Inches(-0.5), Inches(-0.4), Inches(8), Inches(3.2), c(theme, "bokeh", "blush"), 0.5)
    round_rect(slide, Inches(7.5), Inches(-0.2), Inches(7), Inches(2.8), c(theme, "lavender"), 0.5)
    round_rect(slide, Inches(-0.5), Inches(5.6), Inches(14.5), Inches(2.4), c(theme, "peach"), 0.5)
    draw_cloud(slide, 0.3, 1.2, 1.15, c(theme, "cloud"), theme)
    draw_cloud(slide, 9.8, 0.95, 1.0, c(theme, "blush"), theme)
    draw_cloud(slide, 7.0, 5.85, 0.95, c(theme, "cream", "cloud"), theme)
    draw_cloud(slide, 1.4, 6.15, 0.8, c(theme, "cloud"), theme)
    draw_bokeh(slide, theme)


# ---- slides ----

def slide_cover(prs, theme, subtitle=None):
    s = add_blank(prs)
    sky_backdrop(s, theme)
    round_rect(s, Inches(1.2), Inches(1.8), Inches(10.8), Inches(3.8), panel(theme), 0.25)
    draw_poms(s, 1.7, 2.05, theme)
    textbox(s, Inches(1.7), Inches(2.4), Inches(9.8), Inches(0.4),
            "WISHING UMBRELLA MOOD", 14, c(theme, "umbrella", "brand"), True, "Georgia")
    textbox(s, Inches(1.7), Inches(2.9), Inches(8.2), Inches(1.2),
            theme.get("labelJa", theme["label"]), 34, fg(theme), True, "Yu Gothic")
    textbox(s, Inches(1.7), Inches(4.3), Inches(8.0), Inches(0.7),
            subtitle or theme.get("mood", ""), 15, c(theme, "inkSoft", "ink"), False, "Yu Gothic")
    draw_umbrella(s, 10.2, 2.5, 1.05, theme)
    return s


def slide_agenda(prs, theme):
    s = add_blank(prs)
    sky_backdrop(s, theme)
    textbox(s, Inches(0.9), Inches(0.45), Inches(10), Inches(0.5),
            "もくじ", 28, brand(theme), True)
    items = ["雲のかたち", "色", "やわらかUI", "配信・告知", "グッズ", "ありがとう"]
    y = 1.3
    for i, item in enumerate(items, 1):
        round_rect(s, Inches(1.0), Inches(y), Inches(11.2), Inches(0.75), panel(theme), 0.5)
        textbox(s, Inches(1.4), Inches(y + 0.18), Inches(10), Inches(0.45),
                f"{i:02d}  {item}", 18, pfg(theme), True)
        y += 0.9
    return s


def slide_section(prs, theme, num, title, en):
    s = add_blank(prs)
    rect(s, 0, 0, W, H, c(theme, "blush"))
    draw_cloud(s, 0.5, 5.8, 1.3, panel(theme), theme)
    draw_cloud(s, 9.5, 1.0, 1.2, panel(theme), theme)
    round_rect(s, Inches(2.2), Inches(2.2), Inches(8.9), Inches(3.0), panel(theme), 0.3)
    textbox(s, Inches(2.7), Inches(2.5), Inches(8), Inches(0.4), num, 16, brand(theme), True)
    textbox(s, Inches(2.7), Inches(3.1), Inches(8), Inches(0.9), title, 36, pfg(theme), True)
    textbox(s, Inches(2.7), Inches(4.2), Inches(8), Inches(0.4), en, 16, c(theme, "inkSoft", "ink"))
    return s


def slide_about(prs, theme):
    s = add_blank(prs)
    sky_backdrop(s, theme)
    round_rect(s, Inches(0.8), Inches(1.0), Inches(7.4), Inches(5.4), panel(theme), 0.22)
    textbox(s, Inches(1.2), Inches(1.4), Inches(6.6), Inches(0.4), "コンセプト", 14, brand(theme), True)
    textbox(s, Inches(1.2), Inches(2.0), Inches(6.6), Inches(1.0),
            "丸くて、雲みたいな。", 28, pfg(theme), True)
    textbox(s, Inches(1.2), Inches(3.3), Inches(6.6), Inches(2.5),
            "角ばった枠や細い線で真似するより、ふくらんだシルエットとやさしい色面で雰囲気をつくる。",
            16, pfg(theme))
    draw_cloud(s, 8.6, 2.8, 1.4, c(theme, "blush"), theme)
    draw_cloud(s, 9.2, 4.5, 1.0, brand(theme), theme)
    return s


def slide_info(prs, theme):
    s = add_blank(prs)
    sky_backdrop(s, theme)
    textbox(s, Inches(0.9), Inches(0.45), Inches(10), Inches(0.5),
            "おしらせ", 28, brand(theme), True)
    rows = [
        ("TITLE", "まるい雲サンプル"),
        ("DATE", "ふわふわデイ"),
        ("PLACE", "ピンクスカイ"),
        ("STREAM", "ゆるっと配信"),
        ("NOTE", "公式KVは使わず、形と色で寄せる"),
    ]
    y = 1.2
    for label, value in rows:
        round_rect(s, Inches(1.0), Inches(y), Inches(11.2), Inches(0.9), panel(theme), 0.45)
        textbox(s, Inches(1.4), Inches(y + 0.25), Inches(2.2), Inches(0.4), label, 12, brand(theme), True)
        textbox(s, Inches(3.8), Inches(y + 0.22), Inches(7.8), Inches(0.45), value, 18, pfg(theme), True)
        y += 1.05
    return s


def slide_schedule(prs, theme):
    s = add_blank(prs)
    sky_backdrop(s, theme)
    textbox(s, Inches(0.9), Inches(0.45), Inches(10), Inches(0.5),
            "スケジュール", 28, brand(theme), True)
    slots = [("12:00", "準備"), ("15:30", "OPEN"), ("17:00", "START"),
             ("18:30", "休憩雲"), ("19:30", "盛り上がり"), ("20:15", "END")]
    x = 0.7
    for time, label in slots:
        round_rect(s, Inches(x), Inches(2.0), Inches(1.95), Inches(3.6), panel(theme), 0.3)
        oval(s, Inches(x + 0.45), Inches(2.35), Inches(1.05), Inches(1.05), c(theme, "blush"))
        textbox(s, Inches(x + 0.1), Inches(3.6), Inches(1.75), Inches(0.4),
                time, 14, brand(theme), True, align=PP_ALIGN.CENTER)
        textbox(s, Inches(x + 0.1), Inches(4.2), Inches(1.75), Inches(0.8),
                label, 16, pfg(theme), True, align=PP_ALIGN.CENTER)
        x += 2.1
    return s


def slide_cast(prs, theme):
    s = add_blank(prs)
    sky_backdrop(s, theme)
    textbox(s, Inches(0.9), Inches(0.45), Inches(10), Inches(0.5),
            "出演", 28, brand(theme), True)
    round_rect(s, Inches(0.8), Inches(1.3), Inches(5.8), Inches(5.3), panel(theme), 0.22)
    oval(s, Inches(2.1), Inches(1.9), Inches(3.2), Inches(3.2), c(theme, "blush"))
    textbox(s, Inches(1.2), Inches(5.3), Inches(5), Inches(0.7),
            "メイン（まるい枠）", 20, pfg(theme), True, align=PP_ALIGN.CENTER)
    for i, name in enumerate(["ゲストA", "ゲストB", "ゲストC"]):
        top = 1.3 + i * 1.85
        round_rect(s, Inches(7.0), Inches(top), Inches(5.4), Inches(1.65), panel(theme), 0.4)
        oval(s, Inches(7.3), Inches(top + 0.3), Inches(1.05), Inches(1.05), c(theme, "peach"))
        textbox(s, Inches(8.6), Inches(top + 0.55), Inches(3.5), Inches(0.5), name, 18, pfg(theme), True)
    return s


def slide_goods(prs, theme):
    s = add_blank(prs)
    sky_backdrop(s, theme)
    textbox(s, Inches(0.9), Inches(0.45), Inches(10), Inches(0.5),
            "グッズ", 28, brand(theme), True)
    items = ["ぬい", "アクリル", "Tシャツ", "うちわ", "ステッカー", "セット"]
    positions = [(0.8, 1.3), (4.7, 1.3), (8.6, 1.3), (0.8, 4.2), (4.7, 4.2), (8.6, 4.2)]
    for (x, y), name in zip(positions, items):
        round_rect(s, Inches(x), Inches(y), Inches(3.5), Inches(2.5), panel(theme), 0.28)
        oval(s, Inches(x + 0.85), Inches(y + 0.35), Inches(1.8), Inches(1.2), c(theme, "blush"))
        textbox(s, Inches(x + 0.2), Inches(y + 1.75), Inches(3.1), Inches(0.45),
                name, 16, pfg(theme), True, align=PP_ALIGN.CENTER)
    return s


def slide_ticket(prs, theme):
    s = add_blank(prs)
    sky_backdrop(s, theme)
    textbox(s, Inches(0.9), Inches(0.45), Inches(10), Inches(0.5),
            "チケット", 28, brand(theme), True)
    tiers = [("会場", "¥X,XXX"), ("配信", "¥6,500"), ("配信+", "¥7,500")]
    x = 0.9
    for title, price in tiers:
        round_rect(s, Inches(x), Inches(1.5), Inches(3.8), Inches(4.6), panel(theme), 0.25)
        draw_cloud(s, x + 0.55, 2.6, 0.85, c(theme, "blush"), theme)
        textbox(s, Inches(x + 0.3), Inches(3.5), Inches(3.2), Inches(0.4),
                title, 16, brand(theme), True, align=PP_ALIGN.CENTER)
        textbox(s, Inches(x + 0.3), Inches(4.1), Inches(3.2), Inches(0.6),
                price, 28, pfg(theme), True, align=PP_ALIGN.CENTER)
        round_rect(s, Inches(x + 0.55), Inches(5.1), Inches(2.7), Inches(0.65), brand(theme), 0.5)
        textbox(s, Inches(x + 0.55), Inches(5.2), Inches(2.7), Inches(0.45),
                "とる", 16, on_brand(theme), True, align=PP_ALIGN.CENTER)
        x += 4.1
    return s


def slide_streaming(prs, theme):
    s = add_blank(prs)
    sky_backdrop(s, theme)
    round_rect(s, Inches(1.0), Inches(1.1), Inches(11.2), Inches(5.2), panel(theme), 0.2)
    textbox(s, Inches(1.5), Inches(1.5), Inches(10), Inches(0.5), "配信", 24, brand(theme), True)
    for i, line in enumerate([
        "ゆるっと観られる配信枠の案内",
        "アーカイブはふわっと期限まで",
        "チャットもやさしい温度感で",
        "公式素材を使うときは権利確認",
    ]):
        oval(s, Inches(1.7), Inches(2.4 + i * 0.85), Inches(0.45), Inches(0.45), c(theme, "blush"))
        textbox(s, Inches(2.4), Inches(2.35 + i * 0.85), Inches(9), Inches(0.5),
                line, 18, pfg(theme))
    return s


def slide_news(prs, theme):
    s = add_blank(prs)
    sky_backdrop(s, theme)
    textbox(s, Inches(0.9), Inches(0.45), Inches(10), Inches(0.5), "ニュース", 28, brand(theme), True)
    news = [
        ("05.01", "雲グッズ公開"),
        ("04.20", "配信チケ開始"),
        ("04.01", "マップ更新"),
        ("03.15", "ゲスト発表"),
        ("02.01", "特設オープン"),
    ]
    y = 1.25
    for date, title in news:
        round_rect(s, Inches(1.0), Inches(y), Inches(11.2), Inches(0.9), panel(theme), 0.5)
        oval(s, Inches(1.3), Inches(y + 0.2), Inches(0.5), Inches(0.5), brand(theme))
        textbox(s, Inches(2.1), Inches(y + 0.25), Inches(2.2), Inches(0.4), date, 14, brand(theme), True)
        textbox(s, Inches(4.5), Inches(y + 0.22), Inches(7), Inches(0.45), title, 18, pfg(theme), True)
        y += 1.05
    return s


def slide_quote(prs, theme):
    s = add_blank(prs)
    rect(s, 0, 0, W, H, c(theme, "blush"))
    draw_cloud(s, 1.0, 5.5, 1.4, panel(theme), theme)
    draw_cloud(s, 9.0, 1.2, 1.3, panel(theme), theme)
    round_rect(s, Inches(1.8), Inches(2.3), Inches(9.7), Inches(2.8), panel(theme), 0.3)
    textbox(s, Inches(2.3), Inches(2.9), Inches(8.7), Inches(1.2),
            "“ふわっとした一言をここへ”", 28, pfg(theme), True, align=PP_ALIGN.CENTER)
    textbox(s, Inches(2.3), Inches(4.2), Inches(8.7), Inches(0.4),
            "— soft quote", 14, brand(theme), align=PP_ALIGN.CENTER)
    return s


def slide_two_col(prs, theme):
    s = add_blank(prs)
    sky_backdrop(s, theme)
    textbox(s, Inches(0.9), Inches(0.45), Inches(10), Inches(0.5), "ふたつ", 28, brand(theme), True)
    round_rect(s, Inches(0.8), Inches(1.3), Inches(5.7), Inches(5.2), panel(theme), 0.22)
    round_rect(s, Inches(6.8), Inches(1.3), Inches(5.7), Inches(5.2), panel(theme), 0.22)
    textbox(s, Inches(1.2), Inches(1.7), Inches(5), Inches(0.4), "左の雲", 14, brand(theme), True)
    textbox(s, Inches(1.2), Inches(2.4), Inches(5), Inches(3.2),
            "説明をやわらかく。箇条書きより短い段落が似合う。", 16, pfg(theme))
    textbox(s, Inches(7.2), Inches(1.7), Inches(5), Inches(0.4), "右の雲", 14, brand(theme), True)
    textbox(s, Inches(7.2), Inches(2.4), Inches(5), Inches(3.2),
            "FAQやリンクも、角ばった表にせず丸い面に載せる。", 16, pfg(theme))
    return s


def slide_photo(prs, theme):
    s = add_blank(prs)
    sky_backdrop(s, theme)
    round_rect(s, Inches(1.6), Inches(1.0), Inches(10.1), Inches(5.4), panel(theme), 0.2)
    oval(s, Inches(4.5), Inches(2.0), Inches(4.3), Inches(3.2), c(theme, "blush"))
    textbox(s, Inches(1.6), Inches(5.5), Inches(10.1), Inches(0.5),
            "まるいフォト枠 / KEY VISUAL", 18, pfg(theme), True, align=PP_ALIGN.CENTER)
    return s


def slide_access(prs, theme):
    s = add_blank(prs)
    sky_backdrop(s, theme)
    textbox(s, Inches(0.9), Inches(0.45), Inches(10), Inches(0.5), "アクセス", 28, brand(theme), True)
    round_rect(s, Inches(0.8), Inches(1.3), Inches(7.4), Inches(5.2), panel(theme), 0.22)
    textbox(s, Inches(1.2), Inches(1.8), Inches(6.6), Inches(0.5), "会場名", 22, pfg(theme), True)
    textbox(s, Inches(1.2), Inches(2.6), Inches(6.6), Inches(3.0),
            "住所や最寄駅をやさしい文で。\n動線も短く。", 16, pfg(theme))
    round_rect(s, Inches(8.6), Inches(1.3), Inches(3.9), Inches(5.2), c(theme, "peach"), 0.25)
    draw_cloud(s, 9.1, 3.5, 0.95, panel(theme), theme)
    textbox(s, Inches(8.6), Inches(5.4), Inches(3.9), Inches(0.5),
            "MAP", 18, pfg(theme), True, align=PP_ALIGN.CENTER)
    return s


def slide_notes(prs, theme):
    s = add_blank(prs)
    sky_backdrop(s, theme)
    round_rect(s, Inches(1.0), Inches(1.0), Inches(11.2), Inches(5.4), panel(theme), 0.2)
    textbox(s, Inches(1.5), Inches(1.4), Inches(10), Inches(0.5), "注意", 24, brand(theme), True)
    for i, n in enumerate([
        "公式素材の利用条件を確認",
        "角ばったマネより、丸みを優先",
        "1スライドに情報を詰めすぎない",
        "影はふわっと、線は細くしすぎない",
    ]):
        oval(s, Inches(1.7), Inches(2.3 + i * 0.9), Inches(0.5), Inches(0.5), c(theme, "blush"))
        textbox(s, Inches(2.5), Inches(2.3 + i * 0.9), Inches(9), Inches(0.5), n, 18, pfg(theme))
    return s


def slide_sns(prs, theme):
    s = add_blank(prs)
    sky_backdrop(s, theme)
    textbox(s, Inches(0.9), Inches(0.45), Inches(10), Inches(0.5), "リンク", 28, brand(theme), True)
    links = ["X", "YouTube", "BOOTH", "Site"]
    x = 0.9
    for name in links:
        round_rect(s, Inches(x), Inches(2.3), Inches(2.9), Inches(2.6), panel(theme), 0.35)
        oval(s, Inches(x + 0.7), Inches(2.7), Inches(1.5), Inches(1.2), c(theme, "blush"))
        textbox(s, Inches(x), Inches(4.2), Inches(2.9), Inches(0.4),
                name, 16, pfg(theme), True, align=PP_ALIGN.CENTER)
        x += 3.15
    return s


def slide_thanks(prs, theme):
    s = add_blank(prs)
    rect(s, 0, 0, W, H, c(theme, "sky"))
    draw_cloud(s, 1.2, 2.8, 1.6, panel(theme), theme)
    draw_cloud(s, 5.0, 2.2, 1.9, c(theme, "blush"), theme)
    draw_cloud(s, 9.0, 3.0, 1.5, panel(theme), theme)
    textbox(s, Inches(1), Inches(3.0), Inches(11.3), Inches(1.0),
            "ありがとう", 48, pfg(theme), True, align=PP_ALIGN.CENTER)
    textbox(s, Inches(1), Inches(5.5), Inches(11.3), Inches(0.5),
            theme.get("label", ""), 16, brand(theme), align=PP_ALIGN.CENTER)
    return s


def slide_contact(prs, theme):
    s = add_blank(prs)
    sky_backdrop(s, theme)
    round_rect(s, Inches(2.8), Inches(2.0), Inches(7.7), Inches(3.5), panel(theme), 0.28)
    textbox(s, Inches(3.1), Inches(2.5), Inches(7.1), Inches(0.5),
            "コンタクト", 22, brand(theme), True, align=PP_ALIGN.CENTER)
    textbox(s, Inches(3.1), Inches(3.4), Inches(7.1), Inches(1.2),
            "info@example.com\nやさしい文言で案内", 18, pfg(theme), align=PP_ALIGN.CENTER)
    return s


def slide_os_desktop(prs, theme):
    """Uindows-like desktop mock for PPT."""
    s = add_blank(prs)
    # coral desktop wallpaper
    rect(s, 0, 0, W, H, c(theme, "brand", "blush"))
    oval(s, Inches(1.2), Inches(1.0), Inches(4.5), Inches(4.5), c(theme, "cream", "cloud"))
    oval(s, Inches(8.5), Inches(0.4), Inches(3.2), Inches(3.2), c(theme, "paper", "cloud"))
    # menubar
    rect(s, 0, 0, W, Inches(0.45), c(theme, "cream", "paper"))
    textbox(s, Inches(0.3), Inches(0.08), Inches(10), Inches(0.35),
            "File   Edit   View   Tools   Help", 12, c(theme, "ink"), True)
    # icons
    for i, emoji in enumerate(["🎨", "🌂", "🎵", "📅", "🖼", "⏰"]):
        col, row = i % 2, i // 2
        x = 11.2 + col * 0.95
        y = 0.8 + row * 0.95
        round_rect(s, Inches(x), Inches(y), Inches(0.8), Inches(0.8), c(theme, "taupe", "ink"), 0.2)
        textbox(s, Inches(x), Inches(y + 0.2), Inches(0.8), Inches(0.4),
                emoji, 14, c(theme, "white", "cloud"), align=PP_ALIGN.CENTER)
    # window
    round_rect(s, Inches(1.5), Inches(1.3), Inches(8.5), Inches(5.0), c(theme, "paper", "cloud"), 0.12)
    # chunky shadow approximation
    round_rect(s, Inches(1.7), Inches(1.5), Inches(8.5), Inches(5.0), c(theme, "ink"), 0.12)
    round_rect(s, Inches(1.5), Inches(1.3), Inches(8.5), Inches(5.0), c(theme, "paper", "cloud"), 0.12)
    rect(s, Inches(1.5), Inches(1.3), Inches(8.5), Inches(0.55), c(theme, "brand"))
    oval(s, Inches(1.75), Inches(1.45), Inches(0.22), Inches(0.22), c(theme, "cream", "cloud"))
    oval(s, Inches(2.1), Inches(1.45), Inches(0.22), Inches(0.22), c(theme, "cream", "cloud"))
    oval(s, Inches(2.45), Inches(1.45), Inches(0.22), Inches(0.22), c(theme, "cream", "cloud"))
    textbox(s, Inches(1.5), Inches(1.38), Inches(8.5), Inches(0.4),
            "Uindows Desktop", 14, c(theme, "cream", "cloud"), True, align=PP_ALIGN.CENTER)
    textbox(s, Inches(1.9), Inches(2.2), Inches(7.8), Inches(3.5),
            "タイトルバー色分け（red / gray / blue / yellow）\n"
            "ずれた影 10×10\n"
            "デスクトップアイコン 72px角丸\n"
            "コーラルの壁紙＋白い幾何学",
            18, c(theme, "ink"))
    return s


def slide_os_login(prs, theme):
    s = add_blank(prs)
    rect(s, 0, 0, W, H, c(theme, "brand"))
    round_rect(s, Inches(3.5), Inches(2.2), Inches(6.3), Inches(3.2), c(theme, "cream", "cloud"), 0.2)
    textbox(s, Inches(3.5), Inches(2.6), Inches(6.3), Inches(0.6),
            "Uindow OS", 28, c(theme, "brand"), True, align=PP_ALIGN.CENTER)
    textbox(s, Inches(3.5), Inches(3.4), Inches(6.3), Inches(0.5),
            "Login to fluffy desktop", 16, c(theme, "ink"), align=PP_ALIGN.CENTER)
    round_rect(s, Inches(5.1), Inches(4.2), Inches(3.1), Inches(0.65), c(theme, "brand"), 0.5)
    textbox(s, Inches(5.1), Inches(4.3), Inches(3.1), Inches(0.45),
            "ログイン", 16, c(theme, "cream", "cloud"), True, align=PP_ALIGN.CENTER)
    return s


FULL_DECK_BUILDERS = [
    ("cover", lambda prs, t: slide_cover(prs, t, "傘の下のふわふわ秘密基地")),
    ("agenda", slide_agenda),
    ("section-concept", lambda prs, t: slide_section(prs, t, "01", "コンセプト", "WISHING FLUFFY")),
    ("about", slide_about),
    ("section-info", lambda prs, t: slide_section(prs, t, "02", "おしらせ", "INFO")),
    ("info", slide_info),
    ("schedule", slide_schedule),
    ("section-cast", lambda prs, t: slide_section(prs, t, "03", "出演", "CAST")),
    ("cast", slide_cast),
    ("section-goods", lambda prs, t: slide_section(prs, t, "04", "グッズ", "GOODS")),
    ("goods", slide_goods),
    ("section-ticket", lambda prs, t: slide_section(prs, t, "05", "チケット", "TICKET")),
    ("ticket", slide_ticket),
    ("streaming", slide_streaming),
    ("news", slide_news),
    ("quote", slide_quote),
    ("two-col", slide_two_col),
    ("photo", slide_photo),
    ("access", slide_access),
    ("notes", slide_notes),
    ("sns", slide_sns),
    ("os-desktop", slide_os_desktop),
    ("os-login", slide_os_login),
    ("contact", slide_contact),
    ("thanks", slide_thanks),
]


def build_full_deck(theme):
    prs = blank_prs()
    for _, builder in FULL_DECK_BUILDERS:
        builder(prs, theme)
    return prs


def build_pack(themes, title, builders):
    prs = blank_prs()
    s = add_blank(prs)
    rect(s, 0, 0, W, H, hex_rgb("#d9efff"))
    draw_cloud(s, 2.5, 3.2, 1.8, hex_rgb("#ffffff"))
    draw_cloud(s, 7.0, 2.8, 1.6, hex_rgb("#ffd0dc"))
    textbox(s, Inches(1), Inches(3.1), Inches(11.3), Inches(0.9),
            title, 32, hex_rgb("#5a3a44"), True, align=PP_ALIGN.CENTER)
    for theme in themes.values():
        for builder in builders:
            builder(prs, theme)
    return prs


def build_catalog(themes):
    prs = blank_prs()
    s = add_blank(prs)
    rect(s, 0, 0, W, H, hex_rgb("#fff0f4"))
    draw_cloud(s, 3.5, 2.8, 2.0, hex_rgb("#ffffff"))
    textbox(s, Inches(1), Inches(3.0), Inches(11.3), Inches(1),
            "Template Catalog", 40, hex_rgb("#5a3a44"), True, align=PP_ALIGN.CENTER)
    s = add_blank(prs)
    rect(s, 0, 0, W, H, hex_rgb("#fff7fa"))
    textbox(s, Inches(0.8), Inches(0.5), Inches(11), Inches(0.5),
            "LAYOUT INDEX", 28, hex_rgb("#f07a9a"), True)
    for i, (key, _) in enumerate(FULL_DECK_BUILDERS, 1):
        col = 0 if i <= 12 else 1
        row = (i - 1) % 12
        textbox(s, Inches(0.9 + col * 6.2), Inches(1.3 + row * 0.45), Inches(6), Inches(0.4),
                f"{i:02d}  {key}", 14, hex_rgb("#5a3a44"))
    for theme in themes.values():
        slide_cover(prs, theme, "Catalog")
        slide_info(prs, theme)
        slide_ticket(prs, theme)
        slide_thanks(prs, theme)
    return prs


def build_layout_library(themes):
    lib = OUT_DIR / "layouts"
    lib.mkdir(parents=True, exist_ok=True)
    for key, builder in FULL_DECK_BUILDERS:
        prs = blank_prs()
        s = add_blank(prs)
        rect(s, 0, 0, W, H, hex_rgb("#d9efff"))
        draw_cloud(s, 4.5, 3.2, 1.5, hex_rgb("#ffffff"))
        textbox(s, Inches(1), Inches(3.2), Inches(11.3), Inches(0.7),
                f"Layout: {key}", 28, hex_rgb("#5a3a44"), True, align=PP_ALIGN.CENTER)
        for theme in themes.values():
            builder(prs, theme)
        prs.save(lib / f"{key}.pptx")


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    themes = load_themes()
    counts = {}

    for tid, theme in themes.items():
        prs = build_full_deck(theme)
        path = OUT_DIR / f"{tid}-full.pptx"
        prs.save(path)
        counts[path.name] = len(prs.slides)

    packs = {
        "pack-event-announce.pptx": [slide_cover, slide_info, slide_schedule, slide_ticket, slide_sns, slide_thanks],
        "pack-live-day.pptx": [slide_cover, slide_agenda, slide_cast, slide_schedule, slide_goods, slide_notes, slide_thanks],
        "pack-exhibition.pptx": [slide_cover, slide_about, slide_info, slide_goods, slide_access, slide_notes, slide_contact],
        "pack-streaming.pptx": [slide_cover, slide_streaming, slide_ticket, slide_news, slide_sns, slide_thanks],
        "pack-pitch-short.pptx": [slide_cover, slide_about, slide_two_col, slide_quote, slide_ticket, slide_thanks],
        "pack-soft-cloud-only.pptx": [slide_cover, slide_about, slide_info, slide_quote, slide_photo, slide_thanks],
        "pack-uindows-os.pptx": [slide_os_login, slide_os_desktop, slide_about, slide_two_col, slide_sns, slide_thanks],
    }
    for fname, builders in packs.items():
        # soft-cloud-only pack uses only primary theme
        if fname in ("pack-soft-cloud-only.pptx", "pack-uindows-os.pptx"):
            prs = blank_prs()
            s = add_blank(prs)
            rect(s, 0, 0, W, H, hex_rgb("#ff7b87" if "uindows" in fname else "#d9efff"))
            draw_cloud(s, 4, 3, 1.8, hex_rgb("#fff0ed" if "uindows" in fname else "#ffd0dc"))
            title = "Uindows OS Pack" if "uindows" in fname else "Soft Cloud Only"
            textbox(s, Inches(1), Inches(3.1), Inches(11.3), Inches(0.8),
                    title, 32, hex_rgb("#5a3a44"), True, align=PP_ALIGN.CENTER)
            tid = "uindows" if "uindows" in fname else "soft-cloud"
            t = themes[tid]
            for b in builders:
                b(prs, t)
        else:
            prs = build_pack(themes, fname.replace(".pptx", ""), builders)
        path = OUT_DIR / fname
        prs.save(path)
        counts[path.name] = len(prs.slides)

    catalog = build_catalog(themes)
    cpath = OUT_DIR / "catalog-all-themes.pptx"
    catalog.save(cpath)
    counts[cpath.name] = len(catalog.slides)

    build_layout_library(themes)
    layout_count = len(list((OUT_DIR / "layouts").glob("*.pptx")))

    starter = blank_prs()
    s = add_blank(starter)
    rect(s, 0, 0, W, H, hex_rgb("#fff0f4"))
    draw_cloud(s, 3.8, 2.8, 2.0, hex_rgb("#ffffff"))
    textbox(s, Inches(1), Inches(3.0), Inches(11.3), Inches(0.9),
            "Starter Kit — Soft Cloud First", 32, hex_rgb("#5a3a44"), True, align=PP_ALIGN.CENTER)
    for theme in themes.values():
        for _, builder in FULL_DECK_BUILDERS[:8]:
            builder(starter, theme)
    spath = OUT_DIR / "starter-kit-all-themes.pptx"
    starter.save(spath)
    counts[spath.name] = len(starter.slides)

    # update csv with soft-cloud
    lines = ["theme,role,hex"]
    for tid, theme in themes.items():
        for role, hexv in theme["colors"].items():
            lines.append(f"{tid},{role},{hexv}")
    (ROOT / "ppt" / "theme-colors.csv").write_text("\n".join(lines) + "\n", encoding="utf-8")

    print("Generated:")
    for name, n in sorted(counts.items()):
        print(f"  {name}: {n} slides")
    print(f"  layouts/: {layout_count} files")
    print(f"TOTAL pptx: {len(counts) + layout_count}")


if __name__ == "__main__":
    main()
